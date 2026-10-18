"""Génère le modèle financier TotalEnergies (DCF) au format Excel, avec formules vivantes.

Usage : python src/build_model.py [--scenario 1|2|3] [--out chemin.xlsx]
Scénarios : 1 = Bas, 2 = Central, 3 = Haut.
"""
import argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
GREEN = Font(name=F, size=10, color="008000")
GREY = Font(name=F, size=9, italic=True, color="7F7F7F")
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True, color="1F3864")
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
YEL = PatternFill("solid", fgColor="FFFF00")
NUM = '#,##0;(#,##0);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.0x'
PX = '#,##0.00;(#,##0.00);-'


def put(ws, ref, value, font=BLACK, fmt=None, fill=None, align=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(horizontal="center" if i else "left")


def build(path, scenario=2):
    wb = Workbook()

    # ------------------------------------------------------------ Hypotheses
    h = wb.active
    h.title = "Hypotheses"
    put(h, "A1", "TotalEnergies – Hypothèses du modèle DCF", TITLE)
    put(h, "A2", "Légende : bleu = saisie · noir = formule · vert = lien inter-feuilles · fond jaune = hypothèse clé à challenger. Montants en M$ sauf mention.", GREY)

    header(h, 4, ["Paramètres généraux", "Valeur", "Unité", "Source / commentaire"])
    rows = [
        (5, "Scénario retenu (1=Bas, 2=Central, 3=Haut)", scenario, "0", "", "Sélecteur : pilote les lignes « Retenu » ci-dessous"),
        (6, "WACC", 0.09, PCT, "%", "Hypothèse illustrative – à construire (CAPM, bêta, structure cible)"),
        (7, "Croissance terminale (g)", 0.01, PCT, "%", "Hypothèse illustrative – déclin pétrolier long terme vs transition"),
        (8, "Taux d'impôt sur charges financières", 0.35, PCT, "%", "Hypothèse illustrative"),
        (9, "Taux de change (EUR pour 1 USD)", 0.86, '0.000', "EUR/USD", "À METTRE À JOUR avec le spot du jour"),
        (10, "Cours actuel de l'action (EUR)", None, PX, "EUR", "À SAISIR (Euronext Paris : TTE) pour calculer le potentiel"),
    ]
    for r, lab, val, fmt, unit, note in rows:
        put(h, f"A{r}", lab)
        put(h, f"B{r}", val, BLUE, fmt, YEL)
        put(h, f"C{r}", unit)
        put(h, f"D{r}", note, GREY)
    put(h, "C5", '=CHOOSE(B5,"Bas","Central","Haut")', BLACK)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=False)
    h.add_data_validation(dv)
    dv.add("B5")

    header(h, 12, ["Bilan et capital (FY2025)", "Valeur", "Unité", "Source / commentaire"])
    bs = [
        (13, "Capitaux propres", 117523, NUM, "M$", "Form 20-F FY2025 (SEC) – vérifier l'inclusion des minoritaires"),
        (14, "Gearing", 0.15, PCT, "%", "Communiqué résultats FY2025 : gearing ≈ 15 %"),
        (16, "Actions diluées (moyenne pondérée)", 2214, NUM, "M", "Communiqué résultats FY2025"),
        (17, "Trésorerie et équivalents (info)", 26202, NUM, "M$", "Form 20-F FY2025"),
        (18, "Dette financière non courante (info)", 48995, NUM, "M$", "Form 20-F FY2025"),
    ]
    for r, lab, val, fmt, unit, note in bs:
        put(h, f"A{r}", lab)
        put(h, f"B{r}", val, BLUE, fmt)
        put(h, f"C{r}", unit)
        put(h, f"D{r}", note, GREY)
    put(h, "A15", "Dette nette (estimée via le gearing)")
    put(h, "B15", "=B13*B14/(1-B14)", BLACK, NUM)
    put(h, "C15", "M$")
    put(h, "D15", "Estimation : à remplacer par la dette nette publiée", GREY)

    header(h, 21, ["Scénarios opérationnels", "2026E", "2027E", "2028E", "2029E", "2030E"])
    blocks = [
        (22, "Croissance des ventes", PCT, [
            ("Bas", [-0.05, -0.02, 0.0, 0.01, 0.01]),
            ("Central", [0.0, 0.02, 0.02, 0.02, 0.02]),
            ("Haut", [0.03, 0.03, 0.03, 0.025, 0.025])]),
        (27, "Marge CFFO / ventes", PCT, [
            ("Bas", [0.12] * 5),
            ("Central", [0.138] * 5),
            ("Haut", [0.15] * 5)]),
        (32, "Investissements nets (M$)", NUM, [
            ("Bas", [15000] * 5),
            ("Central", [15000] * 5),
            ("Haut", [16000] * 5)]),
    ]
    for r0, title, fmt, cases in blocks:
        put(h, f"A{r0}", title, BOLD)
        for i, (name, vals) in enumerate(cases):
            r = r0 + 1 + i
            put(h, f"A{r}", f"  {name}")
            for j, v in enumerate(vals):
                put(h, f"{'BCDEF'[j]}{r}", v, BLUE, fmt, YEL if name == "Central" else None)
        r = r0 + 4
        put(h, f"A{r}", "  Retenu", BOLD)
        for col in "BCDEF":
            put(h, f"{col}{r}", f"=CHOOSE($B$5,{col}{r0+1},{col}{r0+2},{col}{r0+3})", BLACK, fmt)
    put(h, "A37", "Hypothèses opérationnelles purement illustratives (pas de prévision officielle). Marge CFFO centrale ≈ marge 2025 (27,8 Md$ / 201,2 Md$). "
                  "Investissements nets centraux ≈ 15 Md$/an (ordre de grandeur de la guidance, à vérifier).", GREY)
    h.column_dimensions["A"].width = 46
    for col in "BCDEF":
        h.column_dimensions[col].width = 13

    # ------------------------------------------------------------ Historique
    s = wb.create_sheet("Historique")
    put(s, "A1", "TotalEnergies – Données historiques", TITLE)
    header(s, 4, ["(M$ sauf mention)", "2024", "2025", "Source"])
    hist = [
        (5, "Ventes (Sales)", 214550, 201196, NUM, "Comptes consolidés Q4 2025"),
        (6, "Revenus des ventes", 195610, 182344, NUM, "Comptes consolidés Q4 2025"),
        (7, "Résultat net IFRS (part du Groupe)", 15758, 13127, NUM, "Comptes consolidés Q4 2025"),
        (8, "Résultat net ajusté (part du Groupe)", None, 15587, NUM, "Communiqué FY2025"),
        (9, "CFFO (hors variation de BFR)", None, 27800, NUM, "Communiqué FY2025 : 27,8 Md$ (arrondi)"),
        (10, "Flux de trésorerie d'exploitation (IFRS)", 30854, 27343, NUM, "Form 20-F FY2025"),
        (11, "Investissements nets", None, 17091, NUM, "Communiqué FY2025"),
        (12, "Coût de l'endettement net", 1230, 2067, NUM, "Comptes consolidés Q4 2025"),
        (13, "Actions diluées moyennes (M)", None, 2214, NUM, "Communiqué FY2025"),
        (14, "ROACE", None, 0.126, PCT, "Communiqué FY2025"),
        (15, "Dividende par action (EUR)", None, 3.40, PX, "Communiqué FY2025 (AG du 29/05/2026)"),
    ]
    for r, lab, v24, v25, fmt, src in hist:
        put(s, f"A{r}", lab)
        if v24 is None:
            put(s, f"B{r}", "n.d.", GREY, align="right")
        else:
            put(s, f"B{r}", v24, BLUE, fmt)
        put(s, f"C{r}", v25, BLUE, fmt)
        put(s, f"D{r}", src, GREY)
    put(s, "A17", "Ratios", BOLD)
    put(s, "A18", "Croissance des ventes");        put(s, "C18", "=C5/B5-1", BLACK, PCT)
    put(s, "A19", "Marge CFFO / ventes");          put(s, "C19", "=C9/C5", BLACK, PCT)
    put(s, "A20", "Marge nette ajustée / ventes"); put(s, "C20", "=C8/C5", BLACK, PCT)
    put(s, "A21", "Investissements nets / CFFO");  put(s, "C21", "=C11/C9", BLACK, PCT)
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 12
    s.column_dimensions["C"].width = 12
    s.column_dimensions["D"].width = 48

    # ------------------------------------------------------------ DCF
    d = wb.create_sheet("DCF")
    put(d, "A1", "TotalEnergies – DCF (flux disponibles pour la firme)", TITLE)
    put(d, "A2", '="Scénario : "&Hypotheses!C5&" · valorisation au 31/12/2025 · actualisation fin de période"', GREY)
    header(d, 4, ["(M$)", "2025A", "2026E", "2027E", "2028E", "2029E", "2030E"])
    put(d, "A5", "Période d'actualisation (t)")
    for i, col in enumerate("BCDEFG"):
        put(d, f"{col}5", i, BLUE, "0")
    labels = {6: "Ventes", 7: "  Croissance", 8: "  Marge CFFO / ventes", 9: "CFFO (hors BFR)",
              10: "+ Charges financières nettes après impôt", 11: "– Investissements nets",
              12: "Flux disponible pour la firme (FCFF)", 13: "Facteur d'actualisation", 14: "FCFF actualisé"}
    for r, lab in labels.items():
        put(d, f"A{r}", lab, BOLD if r in (6, 9, 12) else BLACK)
    put(d, "B6", "=Historique!C5", GREEN, NUM)
    put(d, "B8", "=B9/B6", BLACK, PCT)
    put(d, "B9", "=Historique!C9", GREEN, NUM)
    put(d, "B10", "=Historique!C12*(1-Hypotheses!B8)", GREEN, NUM)
    put(d, "B11", "=Historique!C11", GREEN, NUM)
    put(d, "B12", "=B9+B10-B11", BLACK, NUM)
    for i, col in enumerate("CDEFG"):
        prev = "BCDEF"[i]
        hc = "BCDEF"[i]
        put(d, f"{col}6", f"={prev}6*(1+Hypotheses!{hc}26)", GREEN, NUM)
        put(d, f"{col}7", f"={col}6/{prev}6-1", BLACK, PCT)
        put(d, f"{col}8", f"=Hypotheses!{hc}31", GREEN, PCT)
        put(d, f"{col}9", f"={col}6*{col}8", BLACK, NUM)
        put(d, f"{col}10", f"={prev}10", BLACK, NUM)
        put(d, f"{col}11", f"=Hypotheses!{hc}36", GREEN, NUM)
        put(d, f"{col}12", f"={col}9+{col}10-{col}11", BLACK, NUM)
        put(d, f"{col}13", f"=1/(1+Hypotheses!$B$6)^{col}5", GREEN, '0.000')
        put(d, f"{col}14", f"={col}12*{col}13", BLACK, NUM)

    put(d, "A16", "Valorisation", BOLD)
    val = [
        (17, "Somme des FCFF actualisés", "=SUM(C14:G14)", NUM, BLACK),
        (18, "FCFF terminal (2031)", "=G12*(1+Hypotheses!B7)", NUM, GREEN),
        (19, "Valeur terminale", '=IF(Hypotheses!B6>Hypotheses!B7,B18/(Hypotheses!B6-Hypotheses!B7),NA())', NUM, GREEN),
        (20, "Valeur terminale actualisée", "=B19*G13", NUM, BLACK),
        (21, "Valeur d'entreprise (VE)", "=B17+B20", NUM, BLACK),
        (22, "– Dette nette", "=Hypotheses!B15", NUM, GREEN),
        (23, "Valeur des capitaux propres", "=B21-B22", NUM, BLACK),
        (24, "Actions diluées (M)", "=Hypotheses!B16", NUM, GREEN),
        (25, "Valeur par action (USD)", "=B23/B24", PX, BLACK),
        (26, "Valeur par action (EUR)", "=B25*Hypotheses!B9", PX, GREEN),
        (27, "Cours actuel (EUR)", '=IF(ISNUMBER(Hypotheses!B10),Hypotheses!B10,"n.d.")', PX, GREEN),
        (28, "Potentiel / (décote) vs cours", '=IF(ISNUMBER(B27),B26/B27-1,"n.d.")', PCT, BLACK),
        (29, "Poids de la VT dans la VE", "=B20/B21", PCT, BLACK),
        (30, "VE implicite / CFFO 2025", "=B21/B9", MULT, BLACK),
    ]
    for r, lab, f, fmt, fnt in val:
        put(d, f"A{r}", lab, BOLD if r in (21, 26) else BLACK)
        put(d, f"B{r}", f, fnt, fmt, align="right")
    put(d, "A32", "Limites : FCFF approché (CFFO + frais financiers après impôt – investissements nets) ; pas de bridge EBITDA/NOPAT, "
                  "pas de minoritaires ni de provisions de fin de vie. À affiner par segment (E&P, LNG, Power, Downstream).", GREY)
    d.column_dimensions["A"].width = 44
    for col in "BCDEFG":
        d.column_dimensions[col].width = 13

    # ------------------------------------------------------------ Sensibilite
    t = wb.create_sheet("Sensibilite")
    put(t, "A1", "Sensibilité – valeur par action (EUR) : WACC × croissance terminale", TITLE)
    header(t, 4, ["WACC \\ g", "", "", "", "", ""])
    gs = [0.0, 0.005, 0.01, 0.015, 0.02]
    ws_ = [0.07, 0.08, 0.09, 0.10, 0.11]
    for j, g in enumerate(gs):
        c = put(t, f"{'BCDEF'[j]}4", g, BLUE, PCT, align="center")
    for i, w in enumerate(ws_):
        r = 5 + i
        put(t, f"A{r}", w, BLUE, PCT)
        for j in range(5):
            col = "BCDEF"[j]
            f = (f"=(SUMPRODUCT(DCF!$C$12:$G$12/(1+$A{r})^DCF!$C$5:$G$5)"
                 f"+DCF!$G$12*(1+{col}$4)/($A{r}-{col}$4)/(1+$A{r})^DCF!$G$5"
                 f"-Hypotheses!$B$15)/Hypotheses!$B$16*Hypotheses!$B$9")
            put(t, f"{col}{r}", f, GREEN, PX)
    put(t, "A11", "Contrôle vs DCF (0 si WACC/g du modèle sont dans la grille)")
    put(t, "B11", '=IFERROR(ROUND(INDEX(B5:F9,MATCH(Hypotheses!B6,A5:A9,0),MATCH(Hypotheses!B7,B4:F4,0))-DCF!B26,6),"hors grille")', BLACK, '0.000000')
    t.column_dimensions["A"].width = 56
    for col in "BCDEF":
        t.column_dimensions[col].width = 12

    wb.save(path)


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--scenario", type=int, default=2, choices=[1, 2, 3])
    p.add_argument("--out", default="TotalEnergies_DCF.xlsx")
    a = p.parse_args()
    build(a.out, a.scenario)
    print(f"Modèle écrit : {a.out} (scénario {a.scenario})")
